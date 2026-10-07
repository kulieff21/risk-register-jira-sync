"""Read the Excel risk register into typed records that keep their sheet row numbers."""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

SHEET_NAME = "Risks"
REQUIRED_COLUMNS = (
    "risk_id", "title", "asset", "threat", "vulnerability", "owner", "likelihood", "impact",
    "controls", "treatment", "treatment_plan", "due_date", "residual_likelihood",
    "residual_impact", "approved_by",
)


class RegisterError(Exception):
    """The register file cannot be read at all (missing file, no 'Risks' sheet)."""


@dataclass
class Risk:
    row: int
    risk_id: str
    title: str
    asset: str
    threat: str
    vulnerability: str
    owner: str
    likelihood: int | None
    impact: int | None
    controls: list[str]
    treatment: str
    treatment_plan: str
    due_date: dt.date | None
    raw_due_date: str
    residual_likelihood: int | None
    residual_impact: int | None
    approved_by: str

    @property
    def inherent(self) -> int | None:
        if self.likelihood is None or self.impact is None:
            return None
        return self.likelihood * self.impact

    @property
    def residual(self) -> int | None:
        if self.residual_likelihood is None or self.residual_impact is None:
            return None
        return self.residual_likelihood * self.residual_impact


@dataclass
class Register:
    path: Path
    columns: list[str]
    risks: list[Risk] = field(default_factory=list)

    @property
    def missing_columns(self) -> list[str]:
        return [c for c in REQUIRED_COLUMNS if c not in self.columns]


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value).strip()


def parse_int(value) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    s = str(value).strip()
    if re.fullmatch(r"[+-]?\d+(\.0+)?", s):
        return int(float(s))
    return None


def parse_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        try:
            return dt.date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def split_controls(value: str) -> list[str]:
    return [p for p in re.split(r"[;,\s]+", value) if p]


def read_register(path: str | Path) -> Register:
    path = Path(path)
    if not path.is_file():
        raise RegisterError(f"register file not found: {path}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:  # corrupt or not an xlsx
        raise RegisterError(f"cannot read {path}: {exc}") from exc
    if SHEET_NAME not in wb.sheetnames:
        raise RegisterError(f"sheet '{SHEET_NAME}' not found in {path}")
    ws = wb[SHEET_NAME]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    columns = [_text(h).lower() for h in header]
    reg = Register(path=path, columns=columns)
    if reg.missing_columns:
        return reg
    idx = {c: columns.index(c) for c in REQUIRED_COLUMNS}
    for rownum, values in enumerate(rows, start=2):
        if all(v is None or _text(v) == "" for v in values):
            continue

        def cell(name):
            i = idx[name]
            return values[i] if i < len(values) else None

        raw_due = cell("due_date")
        reg.risks.append(Risk(
            row=rownum,
            risk_id=_text(cell("risk_id")),
            title=_text(cell("title")),
            asset=_text(cell("asset")),
            threat=_text(cell("threat")),
            vulnerability=_text(cell("vulnerability")),
            owner=_text(cell("owner")),
            likelihood=parse_int(cell("likelihood")),
            impact=parse_int(cell("impact")),
            controls=split_controls(_text(cell("controls"))),
            treatment=_text(cell("treatment")).lower(),
            treatment_plan=_text(cell("treatment_plan")),
            due_date=parse_date(raw_due),
            raw_due_date=_text(raw_due),
            residual_likelihood=parse_int(cell("residual_likelihood")),
            residual_impact=parse_int(cell("residual_impact")),
            approved_by=_text(cell("approved_by")),
        ))
    return reg
