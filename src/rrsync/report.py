"""Status report (Markdown) and the register copy with Jira status columns (xlsx)."""

from __future__ import annotations

import datetime as dt
from collections import Counter
from pathlib import Path

import openpyxl

from .annex import load_controls
from .plan import group_by_risk, Plan
from .register import SHEET_NAME, Register, Risk

APPENDED = ("jira_key", "jira_status", "resolved", "overdue", "report_date")


def status_of(issue: dict | None) -> tuple[str, str]:
    if issue is None:
        return "not in Jira", ""
    st = issue["fields"].get("status") or {}
    return st.get("name", "unknown"), (st.get("statusCategory") or {}).get("key", "")


def build_rows(risks: list[Risk], issues: list[dict], today: dt.date) -> list[dict]:
    groups = group_by_risk(issues)
    rows = []
    for r in risks:
        found = groups.get(r.risk_id, [])
        issue = found[0] if len(found) == 1 else None
        name, category = status_of(issue)
        if len(found) > 1:
            name = "conflict"
        resolved = ""
        if issue and issue["fields"].get("resolutiondate"):
            resolved = str(issue["fields"]["resolutiondate"])[:10]
        overdue = bool(r.due_date and r.due_date < today and category != "done")
        rows.append({"risk": r, "key": issue["key"] if issue else ("" if not found else ",".join(
            i["key"] for i in found)), "status": name, "category": category,
            "resolved": resolved, "overdue": overdue})
    return rows


def _heat(risks: list[Risk], kind: str) -> list[str]:
    grid = Counter()
    for r in risks:
        l, i = (r.likelihood, r.impact) if kind == "inherent" else (r.residual_likelihood, r.residual_impact)
        grid[(l, i)] += 1
    lines = ["| impact \\ likelihood | 1 | 2 | 3 | 4 | 5 |", "|---|---|---|---|---|---|"]
    for i in range(5, 0, -1):
        cells = [str(grid[(l, i)]) if grid[(l, i)] else "." for l in range(1, 6)]
        lines.append(f"| {i} | " + " | ".join(cells) + " |")
    return lines


def build_markdown(rows: list[dict], plan: Plan, today: dt.date) -> str:
    risks = [r["risk"] for r in rows]
    out = [f"# Risk register status ({today.isoformat()})", "",
           f"Project {plan.project}. {len(risks)} risks in the register.", "",
           "## Treatments", ""]
    for t, n in sorted(Counter(r.treatment for r in risks).items()):
        out.append(f"- {t}: {n}")
    out += ["", "## Jira status", ""]
    for s, n in sorted(Counter(r["status"] for r in rows).items()):
        out.append(f"- {s}: {n}")
    out += ["", "## Heat map: inherent", ""] + _heat(risks, "inherent")
    out += ["", "## Heat map: residual", ""] + _heat(risks, "residual")
    overdue = [r for r in rows if r["overdue"]]
    out += ["", f"## Overdue treatments ({len(overdue)})", ""]
    if overdue:
        out += ["| Risk | Title | Due | Status | Jira |", "|---|---|---|---|---|"]
        for r in sorted(overdue, key=lambda x: (x["risk"].due_date, x["risk"].risk_id)):
            k = r["risk"]
            out.append(f"| {k.risk_id} | {k.title} | {k.due_date} | {r['status']} | {r['key'] or '-'} |")
    else:
        out.append("None.")
    out += ["", "## Orphaned Jira issues", ""]
    orphans = plan.of("orphan")
    out += [f"- {a.key} ({a.risk_id or 'no risk label'})" for a in orphans] or ["None."]
    out += ["", "## Conflicts", ""]
    conflicts = plan.of("conflict")
    out += [f"- {a.risk_id}: {', '.join(a.keys)}" for a in conflicts] or ["None."]
    out += ["", "## Annex A coverage", ""]
    used = {c for r in risks for c in r.controls}
    out += ["| Theme | Referenced | Total |", "|---|---|---|"]
    controls = load_controls()
    for theme in dict.fromkeys(c.theme for c in controls):
        ids = {c.id for c in controls if c.theme == theme}
        out.append(f"| {theme} | {len(ids & used)} | {len(ids)} |")
    total_used = len({c.id for c in controls} & used)
    out.append(f"| all | {total_used} | {len(controls)} |")
    return "\n".join(out) + "\n"


def write_status_xlsx(reg: Register, rows: list[dict], today: dt.date, dest: Path) -> None:
    wb = openpyxl.load_workbook(reg.path)  # a copy in memory; the source file is never saved
    ws = wb[SHEET_NAME]
    start = ws.max_column + 1
    for off, name in enumerate(APPENDED):
        ws.cell(row=1, column=start + off, value=name)
    for r in rows:
        values = (r["key"], r["status"], r["resolved"], "yes" if r["overdue"] else "no", today.isoformat())
        for off, v in enumerate(values):
            ws.cell(row=r["risk"].row, column=start + off, value=v)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
