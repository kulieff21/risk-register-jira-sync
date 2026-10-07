import datetime as dt

import openpyxl

from rrsync.cli import main
from tests.conftest import sha
from tests.fake_jira import FakeJira


def setup_synced(register_path, tmp_path):
    fake = FakeJira()
    out = tmp_path / "out"
    main(["sync", str(register_path), "--project", "RISK", "--apply", "--out", str(out)], client=fake)
    return fake, out


def report(fake, register_path, out, today="2026-10-08"):
    return main(["report", str(register_path), "--project", "RISK", "--out", str(out), "--today", today],
                client=fake)


def status_rows(out):
    ws = openpyxl.load_workbook(out / "register-status.xlsx")["Risks"]
    header = [c.value for c in ws[1]]
    return header, {ws.cell(r, 1).value: dict(zip(header, [c.value for c in ws[r]])) for r in range(2, ws.max_row + 1)}


def test_overdue_logic_uses_status_category(register_path, tmp_path):
    fake, out = setup_synced(register_path, tmp_path)
    k6 = fake.by_label("risk-R-006")[0]
    fake.set_status(k6, "Done", "done", "2026-09-29T10:00:00.000+0000")  # due 09-30, done -> not overdue
    k14 = fake.by_label("risk-R-014")[0]
    fake.set_status(k14, "Done", "indeterminate")  # name says Done but category does not -> still overdue
    k15 = fake.by_label("risk-R-015")[0]
    fake.set_status(k15, "Closed", "done")  # category done with a different status name
    before = sha(register_path)
    fake.writes.clear()
    assert report(fake, register_path, out) == 0
    assert fake.writes == [] and sha(register_path) == before
    header, rows = status_rows(out)
    assert header[-5:] == ["jira_key", "jira_status", "resolved", "overdue", "report_date"]
    assert rows["R-006"]["overdue"] == "no" and rows["R-006"]["resolved"] == "2026-09-29"
    assert rows["R-014"]["overdue"] == "yes"
    assert rows["R-015"]["overdue"] == "no"
    assert rows["R-001"]["overdue"] == "yes" and rows["R-001"]["jira_status"] == "To Do"
    assert rows["R-005"]["overdue"] == "no"  # due 2026-10-31, still in the future
    assert rows["R-006"]["jira_key"] == k6
    assert rows["R-001"]["report_date"] == "2026-10-08"
    md = (out / "report.md").read_text()
    assert "## Overdue treatments (3)" in md  # R-001, R-014, R-022
    assert "R-022" in md and "R-006 |" not in md.split("## Orphaned")[0].split("## Overdue")[1]


def test_today_changes_overdue(register_path, tmp_path):
    fake, out = setup_synced(register_path, tmp_path)
    report(fake, register_path, out, today="2026-08-01")
    _, rows = status_rows(out)
    assert all(r["overdue"] == "no" for r in rows.values())
    report(fake, register_path, out, today="2027-12-31")
    _, rows = status_rows(out)
    dated = [r for r in rows.values() if r["risk_id"] not in ("R-017", "R-023", "R-024")]
    assert all(r["overdue"] == "yes" for r in dated)
    assert rows["R-017"]["overdue"] == "no"  # accepted risks have no due date


def test_markdown_sections(register_path, tmp_path):
    fake, out = setup_synced(register_path, tmp_path)
    fake.add_orphan("R-888")
    fake.add_issue(summary="a", labels=["grc-risk", "risk-R-009"])  # second issue for R-009 -> conflict
    report(fake, register_path, out)
    md = (out / "report.md").read_text()
    for h in ("## Treatments", "## Jira status", "## Heat map: inherent", "## Heat map: residual",
              "## Overdue treatments", "## Orphaned Jira issues", "## Conflicts", "## Annex A coverage"):
        assert h in md
    assert "- mitigate: 19" in md and "- accept: 3" in md and "- transfer: 1" in md and "- avoid: 1" in md
    assert "R-888" in md and "R-009:" in md
    assert "| organizational |" in md and "| technological |" in md
    _, rows = status_rows(out)
    assert rows["R-009"]["jira_status"] == "conflict"


def test_default_today_is_local_date(register_path, tmp_path):
    fake, out = setup_synced(register_path, tmp_path)
    main(["report", str(register_path), "--project", "RISK", "--out", str(out)], client=fake)
    _, rows = status_rows(out)
    assert rows["R-001"]["report_date"] == dt.date.today().isoformat()
