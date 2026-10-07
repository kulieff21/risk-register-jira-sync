import json
import shutil

import openpyxl
import pytest

from rrsync.cli import main
from rrsync.jira import JiraError
from rrsync.plan import build_plan, State
from rrsync.register import read_register
from tests.conftest import EXAMPLES, sha
from tests.fake_jira import FakeJira


def run(fake, reg, out, *extra, cmd="sync"):
    args = [cmd, str(reg), "--project", "RISK", "--out", str(out), *extra]
    return main(args, client=fake)


def audit_lines(out):
    files = sorted((out / "audit").glob("*.jsonl"))
    return [[json.loads(l) for l in f.read_text().splitlines()] for f in files]


def plan_for(fake, reg, out):
    r = read_register(reg)
    return build_plan(r.risks, fake.search_issues("RISK"), State(out / "state.json").project("RISK"), "RISK")


def test_empty_jira_plan_is_all_create_and_dry_run_writes_nothing(register_path, tmp_path, capsys):
    fake = FakeJira()
    out = tmp_path / "out"
    assert run(fake, register_path, out) == 0
    assert fake.writes == []
    assert not (out / "audit").exists()
    assert plan_for(fake, register_path, out).counts()["create"] == 24
    assert "Dry run" in capsys.readouterr().out


def test_apply_then_second_apply_zero_writes(register_path, tmp_path):
    fake = FakeJira()
    out = tmp_path / "out"
    assert run(fake, register_path, out, "--apply") == 0
    assert len(fake.writes) == 24 and len(fake.issues) == 24
    fake.writes.clear()
    assert run(fake, register_path, out, "--apply") == 0
    assert fake.writes == []
    assert plan_for(fake, register_path, out).counts()["noop"] == 24
    # every apply leaves an audit file; the second has only the header
    logs = audit_lines(out)
    assert len(logs) == 2 and len(logs[0]) == 25 and len(logs[1]) == 1
    header = logs[0][0]
    assert header["register_sha256"] == sha(register_path) and header["project"] == "RISK"
    assert header["plan_counts"]["create"] == 24 and header["tool_version"]
    w = logs[0][1]
    assert w["action"] == "create" and w["http_status"] == 201 and w["key"].startswith("RISK-") and w["drift"] is False


def test_created_issue_fields(register_path, tmp_path):
    fake = FakeJira()
    run(fake, register_path, tmp_path / "out", "--apply")
    key = fake.by_label("risk-R-001")[0]
    f = fake.issues[key]["fields"]
    assert f["summary"] == "[R-001] Leaver accounts not disabled"
    assert f["priority"] == {"name": "High"}
    assert f["duedate"] == "2026-09-15"
    assert {"grc-risk", "risk-R-001", "treatment-mitigate", "ctrl-5.18", "ctrl-6.5", "ctrl-5.16"} == set(f["labels"])


def test_lost_create_response_does_not_duplicate(register_path, tmp_path):
    fake = FakeJira()
    out = tmp_path / "out"
    fake.lose_next_create_response = True
    assert run(fake, register_path, out, "--apply") == 2  # one write reported as failed
    assert len(fake.by_label("risk-R-001")) == 1
    assert run(fake, register_path, out, "--apply") == 0
    assert len(fake.issues) == 24
    for i in range(1, 25):
        assert len(fake.by_label(f"risk-R-{i:03d}")) == 1
    fake.writes.clear()
    run(fake, register_path, out, "--apply")
    assert fake.writes == []


def test_hand_edit_is_drift_register_wins_and_audit_records_old_value(register_path, tmp_path, capsys):
    fake = FakeJira()
    out = tmp_path / "out"
    run(fake, register_path, out, "--apply")
    key = fake.by_label("risk-R-005")[0]
    fake.hand_edit(key, summary="something else", priority="Lowest")
    plan = plan_for(fake, register_path, out)
    act = [a for a in plan.actions if a.risk_id == "R-005"][0]
    assert act.kind == "drift" and set(act.drift_fields) == {"summary", "priority"}
    assert plan.counts()["drift"] == 1
    fake.writes.clear()
    capsys.readouterr()
    assert run(fake, register_path, out, "--apply") == 0
    assert "drift in summary" in capsys.readouterr().out
    assert len(fake.writes) == 1
    assert fake.issues[key]["fields"]["summary"] == "[R-005] Phishing leads to credential theft"
    assert fake.issues[key]["fields"]["priority"] == {"name": "Highest"}
    line = audit_lines(out)[-1][1]
    assert line["drift"] is True and line["risk_id"] == "R-005" and line["action"] == "update"
    assert line["changes"]["summary"] == {"before": "something else",
                                          "after": "[R-005] Phishing leads to credential theft"}
    assert line["changes"]["priority"]["before"] == "Lowest"
    # after the overwrite everything is quiet again
    fake.writes.clear()
    run(fake, register_path, out, "--apply")
    assert fake.writes == []


def test_hand_edit_of_description_and_owned_label_is_drift(register_path, tmp_path):
    fake = FakeJira()
    out = tmp_path / "out"
    run(fake, register_path, out, "--apply")
    key = fake.by_label("risk-R-002")[0]
    fake.hand_edit(key, description={"version": 1, "type": "doc", "content": []})
    key2 = fake.by_label("risk-R-003")[0]
    labels = [l for l in fake.issues[key2]["fields"]["labels"] if l != "ctrl-8.2"]
    fake.hand_edit(key2, labels=labels)
    plan = plan_for(fake, register_path, out)
    by = {a.risk_id: a for a in plan.actions}
    assert by["R-002"].kind == "drift" and by["R-002"].drift_fields == ["description"]
    assert by["R-003"].kind == "drift" and by["R-003"].drift_fields == ["labels"]
    run(fake, register_path, out, "--apply")
    assert "ctrl-8.2" in fake.issues[key2]["fields"]["labels"]
    assert fake.issues[key]["fields"]["description"]["content"] != []


def test_hand_added_labels_are_kept_and_not_drift(register_path, tmp_path):
    fake = FakeJira()
    out = tmp_path / "out"
    run(fake, register_path, out, "--apply")
    key = fake.by_label("risk-R-004")[0]
    fake.issues[key]["fields"]["labels"].append("needs-legal")
    assert plan_for(fake, register_path, out).counts()["noop"] == 24
    # a register change on that risk updates labels but keeps the hand-added one
    wb = openpyxl.load_workbook(register_path)
    ws = wb["Risks"]
    ws["I5"] = "5.15, 8.2"  # R-004 controls
    wb.save(register_path)
    plan = plan_for(fake, register_path, out)
    act = [a for a in plan.actions if a.risk_id == "R-004"][0]
    assert act.kind == "update" and act.drift_fields == []
    run(fake, register_path, out, "--apply")
    labels = fake.issues[key]["fields"]["labels"]
    assert "needs-legal" in labels and "ctrl-5.3" not in labels


def test_register_change_is_update_not_drift(register_path, tmp_path):
    fake = FakeJira()
    out = tmp_path / "out"
    run(fake, register_path, out, "--apply")
    wb = openpyxl.load_workbook(register_path)
    wb["Risks"]["B2"] = "Leaver accounts remain active"
    wb.save(register_path)
    fake.writes.clear()
    run(fake, register_path, out, "--apply")
    assert len(fake.writes) == 1
    line = audit_lines(out)[-1][1]
    assert line["drift"] is False and line["changes"]["summary"]["before"] == "[R-001] Leaver accounts not disabled"


def test_duplicate_label_is_conflict_skipped_exit_1(register_path, tmp_path, capsys):
    fake = FakeJira()
    k1 = fake.add_issue(summary="a", labels=["grc-risk", "risk-R-001"])
    k2 = fake.add_issue(summary="b", labels=["grc-risk", "risk-R-001"])
    out = tmp_path / "out"
    assert run(fake, register_path, out, "--apply") == 1
    assert "CONFLICT" in capsys.readouterr().out
    assert [w for w in fake.writes if w[0] == "update"] == []
    assert fake.issues[k1]["fields"]["summary"] == "a" and fake.issues[k2]["fields"]["summary"] == "b"
    assert len(fake.writes) == 23  # all other risks were created
    assert run(fake, register_path, out) == 1  # dry run reports it too


def test_orphan_reported_and_untouched(register_path, tmp_path, capsys):
    fake = FakeJira()
    out = tmp_path / "out"
    run(fake, register_path, out, "--apply")
    key = fake.add_orphan("R-777")
    before = json.dumps(fake.issues[key], sort_keys=True)
    fake.writes.clear()
    capsys.readouterr()
    assert run(fake, register_path, out, "--apply") == 0
    assert "ORPHAN" in capsys.readouterr().out and key in plan_for(fake, register_path, out).of("orphan")[0].key
    assert fake.writes == [] and json.dumps(fake.issues[key], sort_keys=True) == before


def test_issue_with_grc_label_but_no_risk_label_is_orphan(register_path, tmp_path):
    fake = FakeJira()
    k = fake.add_issue(summary="stray", labels=["grc-risk"])
    fake.add_issue(summary="unrelated", labels=["bug"])
    plan = plan_for(fake, register_path, tmp_path / "out")
    assert [a.key for a in plan.of("orphan")] == [k]


def test_validation_errors_block_sync(tmp_path, capsys):
    fake = FakeJira()
    assert main(["sync", str(EXAMPLES / "register-with-errors.xlsx"), "--project", "RISK", "--apply",
                 "--out", str(tmp_path / "out")], client=fake) == 1
    assert fake.writes == [] and fake.reads == 0
    assert "sync refused" in capsys.readouterr().out
    assert not (tmp_path / "out").exists()


def test_update_failure_is_reported_exit_2_and_audited(register_path, tmp_path):
    fake = FakeJira()
    out = tmp_path / "out"
    run(fake, register_path, out, "--apply")
    key = fake.by_label("risk-R-001")[0]
    fake.hand_edit(key, summary="x")
    fake.fail_next_update = JiraError("boom", 400)
    assert run(fake, register_path, out, "--apply") == 2
    line = audit_lines(out)[-1][1]
    assert line["ok"] is False and line["error"] == "boom"


def test_state_file_records_hashes(register_path, tmp_path):
    out = tmp_path / "out"
    run(FakeJira(), register_path, out, "--apply")
    data = json.loads((out / "state.json").read_text())
    entry = data["projects"]["RISK"]["R-001"]
    assert len(entry["desc_hash"]) == 64 and entry["fields"]["priority"] == "High"
    assert len(data["projects"]["RISK"]) == 24
