"""Generate the synthetic example registers (fictional "Larkspur Payments"). Deterministic.

Usage: python tools/make_register.py [output_dir]   (default: examples/)
"""

from __future__ import annotations

import datetime as dt
import io
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

COLUMNS = ["risk_id", "title", "asset", "threat", "vulnerability", "owner", "likelihood", "impact",
           "controls", "treatment", "treatment_plan", "due_date", "residual_likelihood",
           "residual_impact", "approved_by"]

CISO = "CISO Dana Whitcombe <dana.whitcombe@larkspur.example>"
CEO = "CEO Marta Ilves <marta.ilves@larkspur.example>"
ITOPS = "Head of IT Operations"

# (id, title, asset, threat, vulnerability, owner, L, I, controls, treatment, plan, due, rL, rI, approved_by)
GOOD = [
    ("R-001", "Leaver accounts not disabled", "Directory and core payment system accounts",
     "A former employee or contractor signs in with a still-active account",
     "Offboarding is manual and not linked to the HR leaver event", ITOPS, 4, 4,
     "5.18, 6.5, 5.16", "mitigate",
     "Disable accounts automatically from the HR leaver event; weekly reconciliation report",
     "2026-09-15", 2, 3, ""),
    ("R-002", "No periodic user access review", "Core payment system access rights",
     "Access rights accumulate and are never challenged",
     "No scheduled recertification of who has which role", "Compliance Officer", 3, 4,
     "5.15, 5.18, 8.2", "mitigate",
     "Quarterly access recertification signed by each system owner", "2026-11-30", 2, 3, ""),
    ("R-003", "Shared administrator accounts", "Server and database administrator accounts",
     "Administrator actions cannot be attributed to a person",
     "Admin passwords are shared in a team spreadsheet", ITOPS, 4, 4,
     "5.16, 5.17, 8.2", "mitigate",
     "Named admin accounts with a password vault and session logging", "2026-12-15", 2, 3, ""),
    ("R-004", "Excessive privileges in the core payment system", "Core payment system roles",
     "An operator misuses rights to release or alter payments",
     "Roles are broad; operators can both prepare and release payments", "Head of Payments Operations", 3, 5,
     "5.15, 8.2, 5.3", "mitigate",
     "Redesign roles around least privilege and remove unused rights", "2027-01-31", 2, 4, ""),
    ("R-005", "Phishing leads to credential theft", "Staff mailboxes and SSO accounts",
     "Phishing e-mail captures staff credentials",
     "MFA is not enforced everywhere and awareness training is annual only", CISO, 5, 4,
     "6.3, 8.5, 8.7", "mitigate",
     "Phishing-resistant MFA for all staff, quarterly training and simulations", "2026-10-31", 3, 3, ""),
    ("R-006", "Unpatched VPN appliance", "Remote access VPN appliance",
     "Attacker exploits a known vulnerability in the VPN appliance",
     "Firmware is patched quarterly, with no emergency patch process", "Network Engineering Lead", 4, 5,
     "8.8, 8.9, 8.20", "mitigate",
     "Monthly patch window and a 72-hour emergency patch SLA for critical issues", "2026-09-30", 2, 3, ""),
    ("R-007", "Backups never restore-tested", "Payment database backups",
     "Data cannot be recovered after a failure or corruption",
     "Restores have never been tested end to end", ITOPS, 3, 5,
     "8.13, 5.30", "mitigate",
     "Quarterly documented restore test against the payment database", "2026-11-15", 2, 3, ""),
    ("R-008", "Ransomware on file servers and endpoints", "File servers and staff endpoints",
     "Ransomware encrypts data and halts operations",
     "Flat network, uneven endpoint protection, backups reachable from the domain", CISO, 3, 5,
     "8.7, 8.13, 8.22", "mitigate",
     "Network segmentation, EDR on all endpoints, offline backup copy", "2027-02-28", 2, 4, ""),
    ("R-009", "Cloud storage misconfiguration", "Object storage holding settlement files",
     "Settlement files are exposed publicly by a misconfigured bucket",
     "No configuration baseline or automated checks on storage settings", "Cloud Platform Lead", 3, 4,
     "5.23, 8.9", "mitigate",
     "Policy-as-code checks that block public buckets; weekly drift scan", "2026-10-20", 1, 4, ""),
    ("R-010", "Third-party payment processor outage", "Card payment processor connection",
     "The processor is unavailable and payments cannot be authorised",
     "Single processor, no fallback route, weak contractual remedies", "Head of Payments Operations", 3, 4,
     "5.19, 5.22, 5.30", "transfer",
     "Transfer part of the loss through SLA credits and business interruption insurance; agree a fallback processor",
     "2027-01-15", 3, 3, ""),
    ("R-011", "Weak security terms in vendor contracts", "Contracts with ICT suppliers",
     "A supplier incident affects Larkspur without contractual recourse",
     "Older contracts lack security, audit and incident notification clauses", "Head of Legal", 3, 3,
     "5.19, 5.20, 5.22", "mitigate",
     "Standard security addendum; renegotiate the ten most critical suppliers", "2026-12-31", 2, 2, ""),
    ("R-012", "Missing logging and monitoring", "Core payment system and network logs",
     "Malicious activity goes unnoticed for weeks",
     "Logs stay on each host and nobody reviews them", CISO, 4, 4,
     "8.15, 8.16, 8.17", "mitigate",
     "Onboard core systems to a central log platform with alerting and a review rota", "2027-03-15", 2, 3, ""),
    ("R-013", "Insecure code in the merchant portal", "Merchant portal application",
     "Injection or access control flaws expose merchant data",
     "No automated code scanning or security review before release", "Head of Engineering", 4, 4,
     "8.28, 8.25, 8.29", "mitigate",
     "Add SAST and dependency scanning to CI and a pre-release security test", "2026-11-30", 2, 3, ""),
    ("R-014", "Test environments hold real card data", "Test and staging databases",
     "Real card numbers leak from weakly protected test systems",
     "Production data is copied to test for debugging", "Head of Engineering", 3, 5,
     "8.33, 8.11", "avoid",
     "Stop copying production data; use a synthetic data generator and masking", "2026-09-01", 1, 1, ""),
    ("R-015", "Lost or stolen laptops", "Staff laptops",
     "A laptop with company data is lost or stolen",
     "Disk encryption is not enforced on all laptops", "IT Service Desk Manager", 3, 3,
     "8.1, 7.9, 8.24", "mitigate",
     "Enforce full-disk encryption and remote wipe through device management", "2026-10-05", 1, 3, ""),
    ("R-016", "Insider fraud through missing segregation of duties", "Payment initiation and approval workflow",
     "One person creates and approves a fraudulent payment",
     "Initiation and approval can be done by the same user", "Head of Finance", 2, 5,
     "5.3, 5.15", "mitigate",
     "Four-eyes approval for payments above a threshold, enforced in the system", "2026-12-01", 1, 5, ""),
    ("R-017", "Distributed denial of service", "Public payment API and website",
     "A volumetric attack makes the payment API unavailable",
     "Capacity depends on the ISP; no dedicated scrubbing service", CISO, 4, 4,
     "8.20, 8.6", "accept",
     "Rely on ISP-level mitigation for now; revisit at the annual risk review", "", 4, 4, CEO),
    ("R-018", "Weak cryptographic key management", "Encryption keys for card and customer data",
     "Keys are exposed or lost, so data is exposed or unrecoverable",
     "Keys are stored beside the data they protect with no rotation", "Head of Engineering", 2, 5,
     "8.24", "mitigate",
     "Move keys to a managed KMS or HSM with rotation and access logging", "2027-02-15", 1, 4, ""),
    ("R-019", "Personal data kept longer than needed", "Customer personal data stores",
     "Excess personal data is exposed in a breach or breaks data protection rules",
     "No retention schedule and no automated deletion", "Data Protection Officer", 4, 3,
     "5.34, 8.10, 5.33", "mitigate",
     "Approve a retention schedule and automate deletion for the main stores", "2027-01-31", 2, 2, ""),
    ("R-020", "No incident response plan", "Incident handling process",
     "A security incident is handled late or badly",
     "No documented plan, roles or contact tree; no exercises", CISO, 4, 4,
     "5.24, 5.25, 5.26", "mitigate",
     "Write the incident response plan, assign roles and run a tabletop exercise", "2026-10-15", 2, 3, ""),
    ("R-021", "Change management bypassed", "Production change process",
     "An unreviewed change breaks or compromises production",
     "Emergency changes can be deployed without approval or later review", "Head of Engineering", 3, 4,
     "8.32", "mitigate",
     "Require a ticket and a second reviewer for all production changes; review emergency changes weekly",
     "2026-11-20", 2, 3, ""),
    ("R-022", "Uncontrolled access to the server room", "On-premises server room",
     "An unauthorised person gets physical access to servers",
     "A shared door code is used and entries are not logged", "Facilities Manager", 2, 4,
     "7.2, 7.3", "mitigate",
     "Badge access with entry logs and a monthly log review", "2026-08-31", 1, 3, ""),
    ("R-023", "Misuse of privileged utility programs", "Administrative tools on servers",
     "An administrator uses a utility to bypass application controls",
     "Utilities are available to all administrators", ITOPS, 2, 3,
     "8.18", "accept",
     "Utilities limited to the admin group and logged; remaining risk accepted", "", 2, 3, CISO),
    ("R-024", "Clock drift breaks log correlation", "Servers and network devices",
     "Events cannot be correlated across systems during an investigation",
     "Not all devices synchronise to the same time source", "Network Engineering Lead", 2, 2,
     "8.17", "accept",
     "Remaining devices tolerated until the next refresh; remaining risk accepted", "", 2, 2, CISO),
]

# Rows that trigger the row-based validation rules (rule 1, missing columns, is a header-level
# rule and cannot share a file with these rows; the tests build a workbook for it).
BAD = [
    # rule 2: bad id format
    ("RISK-1", "Bad id format", "Asset", "Threat", "Vuln", "Owner", 3, 3, "5.15", "mitigate", "Plan", "2026-12-01", 2, 2, ""),
    # rule 2: duplicate id (twice R-002)
    ("R-002", "Duplicate id, first", "Asset", "Threat", "Vuln", "Owner", 3, 3, "5.15", "mitigate", "Plan", "2026-12-01", 2, 2, ""),
    ("R-002", "Duplicate id, second", "Asset", "Threat", "Vuln", "Owner", 3, 3, "5.15", "mitigate", "Plan", "2026-12-01", 2, 2, ""),
    # rule 3: empty title, owner, asset, threat
    ("R-003", "", "", "", "Vuln", "", 3, 3, "5.15", "mitigate", "Plan", "2026-12-01", 2, 2, ""),
    # rule 4: scores out of range or not integers
    ("R-004", "Bad scores", "Asset", "Threat", "Vuln", "Owner", 7, "high", "5.15", "mitigate", "Plan", "2026-12-01", 0, 2, ""),
    # rule 5: residual greater than inherent
    ("R-005", "Residual above inherent", "Asset", "Threat", "Vuln", "Owner", 2, 2, "5.15", "mitigate", "Plan", "2026-12-01", 4, 4, ""),
    # rule 6: empty controls
    ("R-006", "No controls", "Asset", "Threat", "Vuln", "Owner", 3, 3, "", "mitigate", "Plan", "2026-12-01", 2, 2, ""),
    # rule 6: 2013-style id and an unknown id
    ("R-007", "2013 control id", "Asset", "Threat", "Vuln", "Owner", 3, 3, "A.9.2.3, 5.99", "mitigate", "Plan", "2026-12-01", 2, 2, ""),
    # rule 7: bad treatment
    ("R-008", "Unknown treatment", "Asset", "Threat", "Vuln", "Owner", 3, 3, "5.15", "ignore", "Plan", "2026-12-01", 2, 2, ""),
    # rule 8: missing plan and due date
    ("R-009", "No plan or due date", "Asset", "Threat", "Vuln", "Owner", 3, 3, "5.15", "transfer", "", "", 2, 2, ""),
    # rule 9: accept without approver
    ("R-010", "Accept without approver", "Asset", "Threat", "Vuln", "Owner", 2, 2, "5.15", "accept", "Tolerated", "", 2, 2, ""),
    # rule 10: due date not a date
    ("R-011", "Bad due date", "Asset", "Threat", "Vuln", "Owner", 3, 3, "5.15", "avoid", "Plan", "next Tuesday", 1, 1, ""),
    # warning: high inherent score accepted
    ("R-012", "High score accepted", "Asset", "Threat", "Vuln", "Owner", 5, 4, "5.15", "accept", "Tolerated", "", 5, 4, CISO),
    # warning: mitigate with residual equal to inherent
    ("R-013", "Mitigation changes nothing", "Asset", "Threat", "Vuln", "Owner", 3, 3, "5.15", "mitigate", "Plan", "2026-12-01", 3, 3, ""),
]


def _fix_zip(data: bytes) -> bytes:
    """Make the xlsx byte-identical across runs: fixed zip times and core.xml timestamps."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:(created|modified)[^>]*>)[^<]*", rb"\g<1>2026-10-08T00:00:00Z", content)
            info = zipfile.ZipInfo(item.filename, date_time=(2026, 10, 8, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            dst.writestr(info, content)
    return out.getvalue()


def write_register(rows, path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risks"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        values = list(row)
        due = values[11]
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", due or ""):
            values[11] = dt.date.fromisoformat(due)
        values = [None if v == "" else v for v in values]
        ws.append(values)
    for i, width in enumerate([9, 40, 36, 44, 44, 34, 11, 8, 20, 12, 60, 12, 18, 15, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=12).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "B2"
    wb.properties.creator = "make_register.py"
    wb.properties.created = dt.datetime(2026, 10, 8)
    wb.properties.modified = dt.datetime(2026, 10, 8)
    buf = io.BytesIO()
    wb.save(buf)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_fix_zip(buf.getvalue()))


def main(argv: list[str]) -> None:
    out = Path(argv[0]) if argv else Path(__file__).resolve().parent.parent / "examples"
    write_register(GOOD, out / "register.xlsx")
    write_register(BAD, out / "register-with-errors.xlsx")
    print(f"wrote {out / 'register.xlsx'} and {out / 'register-with-errors.xlsx'}")


if __name__ == "__main__":
    main(sys.argv[1:])
